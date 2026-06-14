"""Document extraction and AI conversion helpers for campaign import."""

from __future__ import annotations

import io
import json
import base64
import multiprocessing
import os
import queue
import re
import shutil
import tempfile
import zipfile
from typing import Any

import fastapi
import json_repair

from . import ai_provider
from .campaign_storage import sanitize_asset_name, unique_path
from .local_config import get_campaign_import_settings


def decode_text_bytes(raw: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "gbk"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise fastapi.HTTPException(status_code=400, detail="文件编码无法识别，请转为 UTF-8 后重试")


def extract_docx_text(raw: bytes) -> str:
    try:
        from docx import Document
    except ImportError:
        # DOCX is a zip of XML files. This fallback loses table structure but
        # keeps import usable when python-docx is missing.
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as zf:
                xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")
            xml = re.sub(r"</w:p\s*>", "\n", xml)
            xml = re.sub(r"<[^>]+>", "", xml)
            return xml
        except Exception as exc:
            raise fastapi.HTTPException(status_code=500, detail="DOCX 解析需要安装 python-docx") from exc

    try:
        doc = Document(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except Exception as exc:
        raise fastapi.HTTPException(status_code=400, detail=f"DOCX 解析失败：{exc}") from exc


def extract_docx_images(raw: bytes, assets_dir: str) -> list[str]:
    saved = []
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            for info in zf.infolist():
                if not info.filename.startswith("word/media/"):
                    continue
                data = zf.read(info.filename)
                if not data:
                    continue
                safe = sanitize_asset_name(os.path.basename(info.filename))
                target = unique_path(assets_dir, safe)
                with open(target, "wb") as f:
                    f.write(data)
                saved.append(os.path.basename(target))
    except Exception:
        pass
    return saved


def _extract_pdf_in_process(
    raw: bytes,
    output_dir: str,
    max_pages: int,
    image_max_pages: int,
    max_images: int,
) -> tuple[str, list[str], list[str]]:
    import pypdf

    warnings: list[str] = []
    saved: list[str] = []
    reader = pypdf.PdfReader(io.BytesIO(raw), strict=False)
    total_pages = len(reader.pages)
    pages_to_read = min(total_pages, max_pages)
    if total_pages > pages_to_read:
        warnings.append(f"PDF 共 {total_pages} 页，仅提取前 {pages_to_read} 页；可在 config.json 调整 campaign_import.pdf_max_pages。")

    text_parts: list[str] = []
    for page_idx in range(pages_to_read):
        try:
            text_parts.append(reader.pages[page_idx].extract_text() or "")
        except Exception as exc:
            warnings.append(f"PDF 第 {page_idx + 1} 页文本提取失败：{type(exc).__name__}")

    image_pages = min(pages_to_read, image_max_pages)
    if image_pages < pages_to_read:
        warnings.append(f"PDF 内嵌图片仅扫描前 {image_pages} 页；可在 config.json 调整 campaign_import.pdf_image_max_pages。")
    for page_idx in range(image_pages):
        if len(saved) >= max_images:
            warnings.append(f"PDF 内嵌图片已达到 {max_images} 个上限；可在 config.json 调整 campaign_import.pdf_max_images。")
            break
        try:
            images = getattr(reader.pages[page_idx], "images", []) or []
        except Exception as exc:
            warnings.append(f"PDF 第 {page_idx + 1} 页图片提取失败：{type(exc).__name__}")
            continue
        for img_idx, image in enumerate(images, start=1):
            if len(saved) >= max_images:
                break
            data = getattr(image, "data", None)
            if not data:
                continue
            raw_name = getattr(image, "name", "") or f"pdf_p{page_idx + 1}_{img_idx}.png"
            safe = sanitize_asset_name(raw_name, ".png")
            target = unique_path(output_dir, safe)
            with open(target, "wb") as f:
                f.write(data)
            saved.append(os.path.basename(target))

    return "\n".join(text_parts), saved, warnings


def _pdf_extract_worker(
    raw: bytes,
    output_dir: str,
    max_pages: int,
    image_max_pages: int,
    max_images: int,
    result_queue: Any,
) -> None:
    try:
        text, saved, warnings = _extract_pdf_in_process(raw, output_dir, max_pages, image_max_pages, max_images)
        result_queue.put({"ok": True, "text": text, "saved": saved, "warnings": warnings})
    except BaseException as exc:
        result_queue.put({
            "ok": False,
            "error_type": type(exc).__name__,
            "message": str(exc) or type(exc).__name__,
        })


def _run_pdf_extract_with_timeout(
    raw: bytes,
    output_dir: str,
    *,
    timeout_seconds: float,
    max_pages: int,
    image_max_pages: int,
    max_images: int,
) -> tuple[str, list[str], list[str]]:
    ctx = multiprocessing.get_context("spawn")
    result_queue = ctx.Queue(maxsize=1)
    process = ctx.Process(
        target=_pdf_extract_worker,
        args=(raw, output_dir, max_pages, image_max_pages, max_images, result_queue),
    )
    process.daemon = True
    process.start()
    process.join(timeout_seconds)
    if process.is_alive():
        process.terminate()
        process.join(3)
        if process.is_alive():
            process.kill()
            process.join(3)
        raise TimeoutError(f"PDF 本地解析超过 {timeout_seconds:g} 秒，已停止本地解析")

    try:
        result = result_queue.get(timeout=2)
    except queue.Empty as exc:
        raise RuntimeError(f"PDF 解析进程无返回，退出码 {process.exitcode}") from exc
    finally:
        result_queue.close()

    if not result.get("ok"):
        raise RuntimeError(f"{result.get('error_type')}: {result.get('message')}")
    return str(result.get("text") or ""), list(result.get("saved") or []), list(result.get("warnings") or [])


def _extract_pdf_text_with_multimodal(raw: bytes) -> tuple[str, list[str]]:
    warnings: list[str] = []
    settings = get_campaign_import_settings()
    page_limit = int(settings.get("pdf_multimodal_pages") or 0)
    if page_limit <= 0:
        return "", warnings
    if not settings["use_ai_conversion"] or not ai_provider.is_configured():
        warnings.append("PDF 未提取到文本；未配置 OpenAI 兼容 Chat 模型，已跳过多模态读取。")
        return "", warnings

    try:
        import fitz
    except ImportError:
        warnings.append("PDF 未提取到文本；如需扫描版 PDF 多模态读取，请安装 PyMuPDF 并选择支持图片输入的 Chat 模型。")
        return "", warnings

    try:
        doc = fitz.open(stream=raw, filetype="pdf")
        pages = min(len(doc), page_limit)
        if pages <= 0:
            return "", warnings
        content: list[dict[str, Any]] = [{
            "type": "text",
            "text": (
                "请读取这些 PDF 页面截图中的中文/英文剧本文字，尽量保持章节、人物、地点、线索和选项结构。"
                "只输出提取到的正文，不要解释。"
            ),
        }]
        for page_idx in range(pages):
            pix = doc.load_page(page_idx).get_pixmap(matrix=fitz.Matrix(1.35, 1.35), alpha=False)
            png = pix.tobytes("png")
            b64 = base64.b64encode(png).decode("ascii")
            content.append({
                "type": "image_url",
                "image_url": {"url": f"data:image/png;base64,{b64}"},
            })
        response = ai_provider.chat_completion(
            [
                {"role": "system", "content": "你是高保真 OCR 与文档转写助手。"},
                {"role": "user", "content": content},
            ],
            temperature=0.0,
            max_tokens=5000,
            timeout=float(settings["ai_timeout_seconds"]),
            apply_token_policy=False,
        )
        text = response.choices[0].message.content or ""
        if text.strip():
            warnings.append(f"本地 PDF 未提取到文本，已使用多模态 Chat 模型读取前 {pages} 页。")
        return text, warnings
    except Exception as exc:
        warnings.append(f"多模态 PDF 读取失败，已继续生成保底剧本：{type(exc).__name__}: {exc}")
        return "", warnings


def extract_pdf(raw: bytes, assets_dir: str) -> tuple[str, list[str], list[str]]:
    warnings: list[str] = []
    saved: list[str] = []
    try:
        import pypdf
    except ImportError as exc:
        raise fastapi.HTTPException(status_code=500, detail="PDF 解析需要安装 pypdf") from exc

    _ = pypdf
    settings = get_campaign_import_settings()
    temp_dir = tempfile.mkdtemp(prefix=".pdf_extract_", dir=assets_dir)
    try:
        try:
            text, temp_saved, pdf_warnings = _run_pdf_extract_with_timeout(
                raw,
                temp_dir,
                timeout_seconds=float(settings["pdf_timeout_seconds"]),
                max_pages=int(settings["pdf_max_pages"]),
                image_max_pages=int(settings["pdf_image_max_pages"]),
                max_images=int(settings["pdf_max_images"]),
            )
            warnings.extend(pdf_warnings)
        except TimeoutError as exc:
            text = ""
            temp_saved = []
            warnings.append(f"{exc}；已继续生成保底剧本。")
        except Exception as exc:
            text = ""
            temp_saved = []
            warnings.append(f"PDF 本地解析失败，已继续生成保底剧本：{exc}")

        for filename in temp_saved:
            source = os.path.join(temp_dir, filename)
            if not os.path.isfile(source):
                continue
            target = unique_path(assets_dir, sanitize_asset_name(filename, ".png"))
            shutil.move(source, target)
            saved.append(os.path.basename(target))
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)

    if not text.strip():
        multimodal_text, multimodal_warnings = _extract_pdf_text_with_multimodal(raw)
        warnings.extend(multimodal_warnings)
        text = multimodal_text.strip()
    if not text.strip():
        warnings.append("PDF 未提取到文本；已生成空白保底剧本。建议另存为 DOCX/TXT，或配置支持图片输入的 Chat 模型并安装 PyMuPDF 后重试扫描版 PDF。")
    return text, saved, warnings


def fallback_campaign(title: str, text: str, first_image_url: str = "") -> dict:
    opening = text.strip()[:1800] or "原始文档未能提取出可用正文。请在导入后手动补充开场场景。"
    return {
        "worldview": f"【{title}】\n\n由原始剧本文档导入生成。AI 转换不可用时创建了保底剧本，请在客户端继续整理节点、地图与触发器。",
        "session_memory": "【跑团记忆日志已初始化】\n",
        "characters": [
            {"name": "玩家", "role": "PC", "hp": 100, "san": 80, "inventory": "", "status": "active"}
        ],
        "nodes": [
            {
                "id": 1,
                "name": "开场",
                "summary": "由导入文档生成的起始场景",
                "content": opening,
                "expanded_content": "",
                "scene_image": first_image_url,
            }
        ],
        "options": [],
        "lorebook": [],
        "triggers": [],
        "world_entities": [],
        "timelines": [],
        "rag_library": [],
        "memory_l1": [],
        "pending_effects": [],
        "npc_chat_logs": [],
    }


def normalize_imported_campaign(raw: dict, title: str, text: str, image_urls: list[str]) -> tuple[dict, dict]:
    campaign = fallback_campaign(title, text, image_urls[0] if image_urls else "")
    if not isinstance(raw, dict):
        return campaign, {"map_rooms": [], "map_edges": []}

    campaign["worldview"] = str(raw.get("worldview") or campaign["worldview"])
    campaign["session_memory"] = str(raw.get("session_memory") or campaign["session_memory"])

    nodes = raw.get("nodes") if isinstance(raw.get("nodes"), list) else []
    normalized_nodes = []
    old_to_new = {}
    for idx, node in enumerate(nodes[:80], start=1):
        if not isinstance(node, dict):
            continue
        old_id = node.get("id", idx)
        old_to_new[str(old_id)] = len(normalized_nodes) + 1
        normalized_nodes.append({
            "id": len(normalized_nodes) + 1,
            "name": str(node.get("name") or f"场景 {idx}")[:80],
            "summary": str(node.get("summary") or "")[:300],
            "content": str(node.get("content") or "")[:6000],
            "expanded_content": str(node.get("expanded_content") or ""),
            "scene_image": str(node.get("scene_image") or ""),
        })
    if normalized_nodes:
        campaign["nodes"] = normalized_nodes
    if image_urls and not any(n.get("scene_image") for n in campaign["nodes"]):
        campaign["nodes"][0]["scene_image"] = image_urls[0]

    valid_node_ids = {n["id"] for n in campaign["nodes"]}
    options = []
    for opt in (raw.get("options") if isinstance(raw.get("options"), list) else [])[:160]:
        if not isinstance(opt, dict):
            continue
        node_id = old_to_new.get(str(opt.get("node_id")), opt.get("node_id"))
        next_id = old_to_new.get(str(opt.get("next_node_id")), opt.get("next_node_id"))
        try:
            node_id = int(node_id)
            next_id = int(next_id)
        except (TypeError, ValueError):
            continue
        if node_id in valid_node_ids and next_id in valid_node_ids:
            options.append({"node_id": node_id, "text": str(opt.get("text") or "继续")[:200], "next_node_id": next_id})
    campaign["options"] = options

    def list_of_dicts(key: str, limit: int) -> list[dict]:
        values = raw.get(key) if isinstance(raw.get(key), list) else []
        return [v for v in values[:limit] if isinstance(v, dict)]

    campaign["characters"] = list_of_dicts("characters", 40) or campaign["characters"]
    campaign["lorebook"] = list_of_dicts("lorebook", 120)
    campaign["triggers"] = list_of_dicts("triggers", 60)
    campaign["world_entities"] = list_of_dicts("world_entities", 80)
    campaign["timelines"] = list_of_dicts("timelines", 20)

    map_rooms = list_of_dicts("map_rooms", 120)
    map_edges = list_of_dicts("map_edges", 240)
    if not map_rooms and isinstance(raw.get("map"), dict):
        map_rooms = [v for v in raw["map"].get("map_rooms", []) if isinstance(v, dict)]
        map_edges = [v for v in raw["map"].get("map_edges", []) if isinstance(v, dict)]
    map_data = {"map_rooms": map_rooms, "map_edges": map_edges}
    return campaign, map_data


def ai_convert_campaign(title: str, text: str, image_urls: list[str]) -> tuple[dict, dict, list[str]]:
    warnings = []
    import_settings = get_campaign_import_settings()
    if not import_settings["use_ai_conversion"]:
        warnings.append("已跳过 AI 剧本转换，生成保底剧本。")
        return fallback_campaign(title, text, image_urls[0] if image_urls else ""), {"map_rooms": [], "map_edges": []}, warnings
    if not ai_provider.is_configured():
        warnings.append("OpenAI 兼容端点尚未配置，已生成保底剧本。")
        return fallback_campaign(title, text, image_urls[0] if image_urls else ""), {"map_rooms": [], "map_edges": []}, warnings

    sample = text[:45000]
    system_prompt = (
        "你是 TRPG 剧本转换器。请把原始剧本文档转换为 Z.R.I.C 客户端可读的 JSON。"
        "只返回 JSON 对象，不要 Markdown。所有内容用中文。"
    )
    user_prompt = f"""
剧本名称：{title}
可用图片 URL：{json.dumps(image_urls, ensure_ascii=False)}

请输出字段：
- worldview: string
- session_memory: string
- characters: array，元素包含 name, role, hp, san, inventory, status
- nodes: array，元素包含 id, name, summary, content, scene_image；id 从 1 开始
- options: array，元素包含 node_id, text, next_node_id
- lorebook: array，元素包含 keywords, content
- triggers: array，可为空
- world_entities: array，元素包含 entity_type, name, location, status, state_desc
- map_rooms: array，元素包含 id, map_id, label, x, y, w, h, description, state, color, node_id, floor
- map_edges: array，元素包含 id, map_id, from_id, to_id, label, locked, key_item, edge_type

要求：
1. 生成一个能直接游玩的起始节点和若干关键场景，优先保持原文结构。
2. 地图房间应绑定相关 node_id；没有把握时少生成，不要编造复杂规则。
3. 图片 URL 只在适合的节点 scene_image 中使用，不要改写 URL。
4. 输出必须是单个 JSON 对象。

原始剧本文档：
{sample}
"""
    try:
        from .agent import _call_ai
        raw = _call_ai(
            system_prompt, user_prompt,
            temperature=0.35, max_tokens=6000, json_mode=True,
            apply_token_policy=False,
            cacheable=False,
            timeout=float(import_settings["ai_timeout_seconds"]),
        )
        if raw.startswith("```"):
            raw = raw.split("```")[1]
            if raw.startswith("json"):
                raw = raw[4:]
        parsed = json_repair.loads(raw.strip())
        campaign, map_data = normalize_imported_campaign(parsed, title, text, image_urls)
        return campaign, map_data, warnings
    except Exception as exc:
        warnings.append(f"AI 转换失败，已生成保底剧本：{type(exc).__name__}: {exc}")
        return fallback_campaign(title, text, image_urls[0] if image_urls else ""), {"map_rooms": [], "map_edges": []}, warnings
